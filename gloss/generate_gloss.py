#!/usr/bin/env python3
"""Генератор базы знаний IPS-MCP из Excel-выгрузок.

Читает All-*-list.xlsx из gloss/, строит SQLite gloss.db и JSON-экспорт
в generated/ по схеме generated/schema.sql.

Запуск:  python gloss/generate_gloss.py
"""

import json
import shutil
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
GENERATED = ROOT / "generated"
BUILD = GENERATED / ".build"

FILES = [
    ("All-objects-list.xlsx", "object_types", "Идентификатор типа объектов"),
    ("All-attributes-list.xlsx", "attribute_types", "Идентификатор атрибута"),
    ("All-links-list.xlsx", "relation_types", "Тип связи"),
]

COLUMNS = {
    "object_types": [
        ("id", "Идентификатор типа объектов", "int"),
        ("name", "Наименование типа объектов", "text"),
        ("object_name", "Наименование объекта", "text"),
        ("versioning", "Версионность", "text"),
        ("comment", "Комментарии", "text"),
        ("default_relation", "Связь по умолчанию", "text"),
        ("guid", "Глобальный идентификатор", "guid"),
        ("domain", "Предметная область", "text"),
        ("descriptor_attribute", "Атрибут-описатель", "text"),
        ("any_attribute", "Возможность присвоения любого атрибута", "text"),
        ("lifecycle", "Жизненный цикл объектов", "text"),
        ("short_name", "Краткое наименование", "text"),
        ("deleted_lifetime", "Время жизни удалённых объектов", "text"),
        ("options", "Опции", "text"),
        ("lifecycle_schema_id", "Идентификатор схемы ЖЦ", "text"),
    ],
    "attribute_types": [
        ("id", "Идентификатор атрибута", "int"),
        ("name", "Наименование", "text"),
        ("short_name", "Краткое наименование", "text"),
        ("alias", "Псевдоним", "text"),
        ("comment", "Комментарии", "text"),
        ("data_type", "Тип данных", "text"),
        ("default_value", "Значение по умолчанию", "text"),
        ("values_list", "Список", "text"),
        ("computed", "Вычисление", "text"),
        ("promotion_level", "Уровень продвижения", "text"),
        ("formula", "Формула", "text"),
        ("language_variant", "Языковой вариант", "text"),
        ("guid", "Глобальный идентификатор", "guid"),
        ("domain", "Предметная область", "text"),
        ("uniqueness", "Уникальность", "text"),
        ("operations_optimization", "Оптимизация операций", "text"),
        ("affects_content_date", "Влияет на дату модификации содержимого объекта", "text"),
        ("options", "Опции", "text"),
        ("input_mask", "Маска ввода значения", "text"),
        ("master_attribute", "Мастер-атрибут", "text"),
        ("data_source", "Источник данных", "text"),
        ("mult_default_values", "F_MULTDEFAULT_VALUES", "text"),
        ("size_type", "Размер/тип", "text"),
    ],
    "relation_types": [
        ("id", "Тип связи", "int"),
        ("name", "Наименование", "text"),
        ("relation_name", "Название связи", "text"),
        ("inverse_relation_name", "Обратное название связи", "text"),
        ("comment", "Комментарий", "text"),
        ("extract_files", "Извлечение файлов объектов", "text"),
        ("relation_kind", "Вид связи", "text"),
        ("guid", "Глобальный идентификатор", "guid"),
        ("domain", "Предметная область", "text"),
        ("any_attribute", "Возможность присвоения любого атрибута", "text"),
        ("short_name", "Краткое наименование", "text"),
        ("options", "Опции", "text"),
    ],
}


class ValidationError(Exception):
    pass


def _cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def read_sheet(path, table):
    """Читает Excel, возвращает (rows, source_modified_at)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    raw = list(ws.iter_rows(values_only=True))
    if len(raw) < 2:
        raise ValidationError(f"{path}: нет строк заголовка")
    hdr = [(_cell(c) if c is not None else "") for c in raw[1]]
    colmap = COLUMNS[table]

    missing = [t for _, t, _ in colmap if t not in hdr]
    if missing:
        raise ValidationError(f"{path}: отсутствуют колонки {missing}")

    idx = {t: hdr.index(t) for _, t, _ in colmap}

    rows = []
    for r in raw[2:]:
        rec = {}
        for field, title, kind in colmap:
            v = _cell(r[idx[title]])
            if kind == "int":
                if v is None:
                    continue
                try:
                    rec[field] = int(v)
                except (TypeError, ValueError):
                    raise ValidationError(f"{path}: нечисловой ID {v!r}")
            else:
                rec[field] = v
        rows.append(rec)
    mtime = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat()
    return rows, mtime


def validate(rows, table, errors):
    ids = set()
    for r in rows:
        rid = r.get("id")
        if rid is None:
            errors.append((table, None, "missing_id", "пустой ID"))
            continue
        if rid in ids:
            errors.append((table, rid, "duplicate_id", f"дубликат ID {rid}"))
        ids.add(rid)
        if not r.get("name"):
            errors.append((table, rid, "empty_name", "пустое имя"))
        guid = r.get("guid")
        if guid and (len(guid) != 36 or guid.count("-") != 4):
            errors.append((table, rid, "invalid_guid", f"некорректный GUID {guid!r}"))


def build_json(rows):
    return [{"id": r["id"], "name": r["name"], "guid": r.get("guid")} for r in rows]


def build_json_by_id(rows):
    return {str(r["id"]): {"id": r["id"], "name": r["name"], "guid": r.get("guid")} for r in rows}


def main():
    for f, table, _ in FILES:
        src = ROOT / f
        if not src.exists():
            raise ValidationError(f"отсутствует исходный файл {src}")

    if BUILD.exists():
        shutil.rmtree(BUILD)
    BUILD.mkdir(parents=True)

    errors = []
    data = {}
    source_files = []
    for f, table, _ in FILES:
        src = ROOT / f
        rows, mtime = read_sheet(src, table)
        validate(rows, table, errors)
        data[table] = rows
        source_files.append({"file": f, "modified_at": mtime, "table": table})

    if errors:
        for e in errors:
            print(f"ERROR: {e}", file=sys.stderr)
        raise ValidationError(f"{len(errors)} ошибок валидации")

    db_path = BUILD / "gloss.db"
    conn = sqlite3.connect(db_path)
    schema = (GENERATED / "schema.sql").read_text(encoding="utf-8")
    conn.executescript(schema)

    colnames = {t: [c[0] for c in COLUMNS[t]] + [] for t in COLUMNS}

    for f, table, _ in FILES:
        src = ROOT / f
        cols = [c[0] for c in COLUMNS[table]] + ["source_file", "source_modified_at"]
        for r in data[table]:
            values = []
            for c in COLUMNS[table]:
                values.append(r.get(c[0]))
            values.append(f)
            values.append(source_files[[s["table"] for s in source_files].index(table)]["modified_at"])
            placeholders = ",".join("?" for _ in cols)
            conn.execute(
                f"INSERT INTO {table} ({','.join(cols)}) VALUES ({placeholders})",
                values,
            )
    conn.commit()

    generated_at = datetime.now(timezone.utc).isoformat()
    counts = {t: len(data[t]) for t in sorted(data)}

    with open(BUILD / "object-types.json", "w", encoding="utf-8") as fh:
        json.dump(build_json(data["object_types"]), fh, ensure_ascii=False, indent=2)
    with open(BUILD / "attribute-types.json", "w", encoding="utf-8") as fh:
        json.dump(build_json(data["attribute_types"]), fh, ensure_ascii=False, indent=2)
    with open(BUILD / "relation-types.json", "w", encoding="utf-8") as fh:
        json.dump(build_json(data["relation_types"]), fh, ensure_ascii=False, indent=2)
    with open(BUILD / "object-types.by-id.json", "w", encoding="utf-8") as fh:
        json.dump(build_json_by_id(data["object_types"]), fh, ensure_ascii=False, indent=2)
    with open(BUILD / "attribute-types.by-id.json", "w", encoding="utf-8") as fh:
        json.dump(build_json_by_id(data["attribute_types"]), fh, ensure_ascii=False, indent=2)
    with open(BUILD / "relation-types.by-id.json", "w", encoding="utf-8") as fh:
        json.dump(build_json_by_id(data["relation_types"]), fh, ensure_ascii=False, indent=2)

    manifest = {
        "schemaVersion": 1,
        "generatedAt": generated_at,
        "sourceFiles": source_files,
        "counts": {
            "objectTypes": counts["object_types"],
            "attributeTypes": counts["attribute_types"],
            "relationTypes": counts["relation_types"],
        },
        "databaseId": None,
        "ipsVersion": None,
        "metadataGeneration": None,
        "validation": {
            "duplicateIds": 0,
            "invalidGuids": 0,
            "emptyNames": 0,
        },
    }
    with open(BUILD / "GENERATE-MANIFEST.json", "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=2)

    smoke(conn)
    conn.close()

    for p in [
        "gloss.db", "object-types.json", "attribute-types.json", "relation-types.json",
        "object-types.by-id.json", "attribute-types.by-id.json", "relation-types.by-id.json",
        "GENERATE-MANIFEST.json",
    ]:
        shutil.copy2(BUILD / p, GENERATED / p)

    print(f"OK: {counts['object_types']} object types, "
          f"{counts['attribute_types']} attributes, {counts['relation_types']} relations")
    print(f"    -> {GENERATED / 'gloss.db'}")


def smoke(conn):
    checks = [
        ("SELECT COUNT(*) FROM object_types", None),
        ("SELECT COUNT(*) FROM attribute_types", None),
        ("SELECT COUNT(*) FROM relation_types", None),
        ("SELECT id, name FROM relation_types WHERE id = 1002", 1002),
        ("SELECT id, name FROM attribute_types WHERE id IN (9,10,1066,1032) ORDER BY id", 1066),
        ("SELECT id, name FROM object_types WHERE id IN (1075,1110,1212) ORDER BY id", 1075),
    ]
    for q, expect in checks:
        row = conn.execute(q).fetchone()
        if row is None:
            raise ValidationError(f"smoke-проверка не пройдена: {q}")
        print(f"  smoke: {row}")


if __name__ == "__main__":
    try:
        main()
    except ValidationError as e:
        print(f"FAILED: {e}", file=sys.stderr)
        sys.exit(1)
